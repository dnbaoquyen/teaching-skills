"""MKT1107 final oral exam (group) grading workbook. usage: python3 build_oral.py <out.xlsx>"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]
F = "Arial"
NG, NS = 25, 60  # group rows, student rows

bold = Font(name=F, bold=True)
norm = Font(name=F)
inp = Font(name=F, color="0000FF")
title = Font(name=F, bold=True, size=13, color="1F4E79")
head_fill = PatternFill("solid", fgColor="DCE6F1")
in_fill = PatternFill("solid", fgColor="FFF9E5")
out_fill = PatternFill("solid", fgColor="EAF1DD")
thin = Side(style="thin", color="A6A6A6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap = Alignment(wrap_text=True, vertical="top")


def hdr(ws, row, vals, widths=None, start=1):
    for j, v in enumerate(vals, start):
        c = ws.cell(row, j, v)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, center, box
    if widths:
        for j, w in enumerate(widths, start):
            ws.column_dimensions[L(j)].width = w


def cell(ws, r, c, v, kind="f", fmt=None, align=center):
    x = ws.cell(r, c, v)
    x.border = box
    x.alignment = align
    if kind == "in":
        x.font, x.fill = inp, in_fill
    elif kind == "out":
        x.font, x.fill = bold, out_fill
    else:
        x.font = norm
    if fmt:
        x.number_format = fmt
    return x


wb = Workbook()

# ================= Hướng dẫn =================
hd = wb.active
hd.title = "Hướng dẫn"
hd.column_dimensions["A"].width = 4
hd.column_dimensions["B"].width = 110
hd["B1"] = "MKT1107 · Bảng chấm thi vấn đáp cuối kỳ theo nhóm"; hd["B1"].font = title
lines = [
    "Giảng viên: Đoàn Nguyễn Bảo Quyên · Thang 10 · Rubric chi tiết: tài liệu “MKT1107 · Rubric thi vấn đáp cuối kỳ theo nhóm”.",
    "",
    "CÁCH DÙNG",
    "1. Sheet “Chấm nhóm”: mỗi nhóm một dòng. Nhập mã nhóm (ví dụ N01), đề tài, hướng, rồi chọn điểm từng tiêu chí A1–A5, B1–B4, C1, C2 từ danh sách thả xuống.",
    "2. Nếu một câu chung được 0 hoặc 0,5 điểm, cột “Được hỏi vớt?” báo “Có”. Hỏi nhóm có nhận câu vớt không, rồi chọn ở cột “Câu vớt”: Không cần / Nhóm từ chối / Đã hỏi.",
    "3. Nếu “Đã hỏi”: nhập điểm câu vớt (0–2). Điểm câu vớt tự thay cho câu chung không đạt (câu thấp hơn nếu cả hai không đạt), nếu cao hơn.",
    "4. Sheet “Chấm cá nhân”: mỗi sinh viên một dòng. Chọn mã nhóm, nhập điểm câu cá nhân D (0–1). Điểm cuối, xếp loại tự tính.",
    "5. Sheet “Tổng hợp”: thống kê điểm và xếp loại của lớp. Sheet “Ví dụ”: 4 ví dụ trong rubric, tính bằng cùng công thức để đối chiếu.",
    "",
    "MÀU Ô",
    "Chữ xanh trên nền vàng nhạt = ô nhập. Nền xanh lá nhạt = kết quả tự tính (không sửa). Ô đỏ = cảnh báo cần xem lại.",
    "",
    "QUY TẮC TÍNH ĐIỂM (theo rubric)",
    "Điểm nhóm = A (3) + B (2) + C1 + C2 (4), sau khi thay điểm câu vớt nếu có. Điểm sinh viên = Điểm nhóm + D (1). Làm tròn đến 0,25.",
    "Nhóm đã nhận câu vớt: điểm cuối tối đa 8. Cách áp trần chọn ở ô C4 sheet “Chấm cá nhân”:",
    "   • “Điểm cuối” (mặc định): Điểm cuối = MIN(8; Điểm nhóm + D).",
    "   • “Phần nhóm”: Điểm cuối = MIN(7; Điểm nhóm) + D — giữ khác biệt giữa các thành viên.",
    "Xếp loại theo phụ lục đề cương: Xuất sắc 9–10 · Giỏi 8–<9 · Khá 6,5–<8 · Trung bình 5–<6,5 · Yếu 3–<5 · Kém 0–<3.",
    "",
    "LƯU Ý",
    "File chứa họ tên, MSSV sinh viên khi đã nhập: lưu và chia sẻ theo quy định bảo mật dữ liệu của trường.",
    "Sinh viên vắng thi, thành viên không đóng góp: xử lý theo quyết định của giảng viên và quy chế trường, ghi ở cột Ghi chú.",
]
for i, t in enumerate(lines, 2):
    c = hd.cell(i, 2, t)
    c.font = bold if t.isupper() and t else norm
    c.alignment = Alignment(wrap_text=True, vertical="top")

# ================= Rubric =================
rb = wb.create_sheet("Rubric")
hdr(rb, 1, ["Mã", "Tiêu chí", "Tối đa", "Mức điểm cho phép", "Mô tả mức cao nhất"], [7, 42, 9, 22, 80])
crit = [
    ("A1", "Vấn đề, mục tiêu, câu hỏi nghiên cứu (Bài 1, 4)", 0.5, "0 · 0,25 · 0,5", "VĐ quản trị/nghiên cứu tách bạch; mục tiêu SMART; câu hỏi khớp mục tiêu; đối tượng, phạm vi rõ"),
    ("A2", "Cơ sở lý thuyết, tổng quan, trích dẫn (Bài 2, 3)", 0.5, "0 · 0,25 · 0,5", "Tổng hợp nguồn học thuật có thật → mô hình/khung chủ đề; APA 7 đúng, khớp hai chiều; khai báo và nhật ký AI"),
    ("A3", "Phương pháp nghiên cứu (Bài 5–8)", 0.75, "0 · 0,25 · 0,5 · 0,75", "Phương pháp hợp câu hỏi; đúng tên phương pháp chọn mẫu, cỡ mẫu có căn cứ; thang đo có nguồn, đúng cấp; có thử nghiệm; tuân thủ đạo đức"),
    ("A4", "Kết quả và phân tích (Bài 9)", 0.75, "0 · 0,25 · 0,5 · 0,75", "Báo cáo số phiếu; kỹ thuật đúng cấp thang đo, trả lời câu hỏi NC; diễn giải đúng (bác bỏ/chưa đủ cơ sở bác bỏ H0; tương quan ≠ nhân quả)"),
    ("A5", "Kết luận, hàm ý quản trị, hạn chế (Bài 4, 10)", 0.5, "0 · 0,25 · 0,5", "Trả lời đủ câu hỏi; hàm ý cụ thể bám kết quả; nêu trung thực hạn chế khái quát hóa"),
    ("B1", "Cấu trúc và thời gian (Bài 10)", 0.5, "0 · 0,25 · 0,5", "Vấn đề → phương pháp → kết quả → hàm ý; xong trong 3 phút"),
    ("B2", "Bảng, biểu đồ (Bài 10 mục 10.3)", 0.5, "0 · 0,25 · 0,5", "Có tên, đơn vị, nguồn; đúng loại biểu đồ; không làm sai lệch cảm nhận"),
    ("B3", "Rõ ràng, chính xác thuật ngữ", 0.5, "0 · 0,25 · 0,5", "Nêu rõ phát hiện chính bằng số liệu; dùng đúng thuật ngữ"),
    ("B4", "Phối hợp nhóm (AM9)", 0.5, "0 · 0,25 · 0,5", "Các thành viên cùng trình bày/hỗ trợ rõ vai trò"),
    ("C1", "Câu hỏi chung 1", 2, "0 · 0,5 · 1 · 1,5 · 2", "Đúng, đủ ý, giải thích được, liên hệ cụ thể đề tài và số liệu nhóm"),
    ("C2", "Câu hỏi chung 2", 2, "0 · 0,5 · 1 · 1,5 · 2", "Như C1"),
    ("V", "Câu hỏi vớt (nếu có)", 2, "0 · 0,5 · 1 · 1,5 · 2", "Như C1; thay cho câu chung không đạt; nhóm đã nhận vớt thì điểm cuối tối đa 8"),
    ("D", "Câu hỏi cá nhân", 1, "0 · 0,25 · 0,5 · 0,75 · 1", "Đúng, đủ, giải thích được và ví dụ từ dự án nhóm"),
]
for i, row in enumerate(crit, 2):
    for j, v in enumerate(row, 1):
        cell(rb, i, j, v, align=center if j in (1, 3, 4) else wrap)
r = len(crit) + 2
cell(rb, r, 2, "Tổng tối đa", align=wrap).font = bold
cell(rb, r, 3, "=SUM(C2:C12)+C14", kind="out")

# ================= Chấm nhóm =================
g = wb.create_sheet("Chấm nhóm")
g["A1"] = "CHẤM PHẦN NHÓM"; g["A1"].font = title
g["A2"] = "Chọn điểm từ danh sách thả xuống. Cột xanh lá tự tính."; g["A2"].font = Font(name=F, italic=True)
cols = ["Mã nhóm", "Đề tài", "Hướng", "A1\n(0,5)", "A2\n(0,5)", "A3\n(0,75)", "A4\n(0,75)", "A5\n(0,5)",
        "B1\n(0,5)", "B2\n(0,5)", "B3\n(0,5)", "B4\n(0,5)", "C1\n(2)", "C2\n(2)",
        "Được hỏi vớt?", "Câu vớt", "Điểm vớt\n(2)", "Vớt thay cho",
        "A\n(3)", "B\n(2)", "C sau vớt\n(4)", "Điểm nhóm\n(9)", "Nhóm đã nhận vớt", "Kiểm tra", "Ghi chú"]
widths = [9, 30, 11] + [7] * 11 + [10, 14, 9, 10, 8, 8, 10, 11, 10, 22, 30]
HR = 4
hdr(g, HR, cols, widths)
g.row_dimensions[HR].height = 32
g.freeze_panes = "D5"
first, last = HR + 1, HR + NG
dv = {
    "half": DataValidation(type="list", formula1='"0,0.25,0.5"', allow_blank=True),
    "q": DataValidation(type="list", formula1='"0,0.25,0.5,0.75"', allow_blank=True),
    "two": DataValidation(type="list", formula1='"0,0.5,1,1.5,2"', allow_blank=True),
    "dir": DataValidation(type="list", formula1='"Định lượng,Định tính"', allow_blank=True),
    "vot": DataValidation(type="list", formula1='"Không cần,Nhóm từ chối,Đã hỏi"', allow_blank=True),
}
for d in dv.values():
    d.error = "Chọn giá trị trong danh sách."; d.showErrorMessage = True
    g.add_data_validation(d)
for r in range(first, last + 1):
    for c in range(1, 4):
        cell(g, r, c, None, "in", align=wrap if c == 2 else center)
    for c in range(4, 15):
        cell(g, r, c, None, "in", fmt="0.00")
    dv["dir"].add(f"C{r}")
    for c in (4, 5, 8, 9, 10, 11, 12):
        dv["half"].add(f"{L(c)}{r}")
    for c in (6, 7):
        dv["q"].add(f"{L(c)}{r}")
    for c in (13, 14):
        dv["two"].add(f"{L(c)}{r}")
    # O: eligible for rescue question
    cell(g, r, 15, f'=IF(COUNT(M{r}:N{r})=0,"",IF(OR(AND(ISNUMBER(M{r}),M{r}<=0.5),AND(ISNUMBER(N{r}),N{r}<=0.5)),"Có","Không"))', "out")
    cell(g, r, 16, None, "in"); dv["vot"].add(f"P{r}")
    cell(g, r, 17, None, "in", fmt="0.00"); dv["two"].add(f"Q{r}")
    # R: which common question the rescue replaces (the lower one)
    cell(g, r, 18, f'=IF(P{r}="Đã hỏi",IF(MIN(N(M{r}),N(N{r}))<=0.5,IF(N(M{r})<=N(N{r}),"C1","C2"),""),"")', "out")
    cell(g, r, 19, f'=IF(COUNT(D{r}:H{r})=0,"",SUM(D{r}:H{r}))', "out", fmt="0.00")
    cell(g, r, 20, f'=IF(COUNT(I{r}:L{r})=0,"",SUM(I{r}:L{r}))', "out", fmt="0.00")
    cell(g, r, 21, f'=IF(COUNT(M{r}:N{r})=0,"",IF(R{r}="C1",MAX(N(M{r}),N(Q{r})),N(M{r}))+IF(R{r}="C2",MAX(N(N{r}),N(Q{r})),N(N{r})))', "out", fmt="0.00")
    cell(g, r, 22, f'=IF(A{r}="","",N(S{r})+N(T{r})+N(U{r}))', "out", fmt="0.00")
    cell(g, r, 23, f'=IF(A{r}="","",IF(P{r}="Đã hỏi","Có","Không"))', "out")
    cell(g, r, 24,
         f'=IF(A{r}="","",IF(COUNT(D{r}:N{r})<11,"Chưa chấm đủ tiêu chí",'
         f'IF(AND(P{r}="Đã hỏi",O{r}<>"Có"),"Nhóm không đủ điều kiện hỏi vớt",'
         f'IF(AND(P{r}="Đã hỏi",Q{r}=""),"Chưa nhập điểm vớt",'
         f'IF(AND(O{r}="Có",P{r}=""),"Chọn trạng thái câu vớt","OK")))))', "out", align=wrap)
    cell(g, r, 25, None, "in", align=wrap)
g.conditional_formatting.add(f"X{first}:X{last}", FormulaRule(formula=[f'AND(X{first}<>"",X{first}<>"OK")'],
                             fill=PatternFill("solid", fgColor="F8CBAD")))

# ================= Chấm cá nhân =================
s = wb.create_sheet("Chấm cá nhân")
s["A1"] = "CHẤM PHẦN CÁ NHÂN VÀ ĐIỂM CUỐI"; s["A1"].font = title
s["A2"] = "Chọn mã nhóm (phải có ở sheet “Chấm nhóm”), nhập điểm câu cá nhân D."; s["A2"].font = Font(name=F, italic=True)
s["A4"] = "Cách áp trần 8 khi nhóm nhận câu vớt:"; s["A4"].font = bold
cap = s["C4"]; cap.value = "Điểm cuối"; cap.font, cap.fill, cap.border = inp, in_fill, box
capdv = DataValidation(type="list", formula1='"Điểm cuối,Phần nhóm"', allow_blank=False)
s.add_data_validation(capdv); capdv.add("C4")
cap.comment = Comment("Điểm cuối: MIN(8; nhóm + D). Phần nhóm: MIN(7; nhóm) + D.", "MKT1107")
SH = 6
hdr(s, SH, ["STT", "MSSV", "Họ và tên", "Mã nhóm", "Câu cá nhân (mã)", "D\n(1)", "Điểm nhóm\n(9)",
            "Nhóm đã nhận vớt", "Điểm trước trần", "Điểm cuối\n(10)", "Xếp loại", "Kiểm tra", "Ghi chú"],
    [6, 13, 26, 10, 12, 8, 11, 11, 11, 11, 12, 22, 30])
s.row_dimensions[SH].height = 32
s.freeze_panes = "D7"
gdv = DataValidation(type="list", formula1=f"='Chấm nhóm'!$A${first}:$A${last}", allow_blank=True)
ddv = DataValidation(type="list", formula1='"0,0.25,0.5,0.75,1"', allow_blank=True)
for d in (gdv, ddv):
    d.error = "Chọn giá trị trong danh sách."; d.showErrorMessage = True
    s.add_data_validation(d)
GA, GV, GW = (f"'Chấm nhóm'!$A${first}:$A${last}", f"'Chấm nhóm'!$V${first}:$V${last}",
              f"'Chấm nhóm'!$W${first}:$W${last}")
s1, s2 = SH + 1, SH + NS
for i in range(NS):
    r = s1 + i
    cell(s, r, 1, f'=IF(B{r}="","",ROW()-{SH})', "f")
    for c in (2, 3, 4, 5):
        cell(s, r, c, None, "in", align=wrap if c == 3 else center)
    gdv.add(f"D{r}")
    cell(s, r, 6, None, "in", fmt="0.00"); ddv.add(f"F{r}")
    m = f"MATCH(D{r},{GA},0)"
    cell(s, r, 7, f'=IF(D{r}="","",IFERROR(INDEX({GV},{m}),""))', "out", fmt="0.00")
    cell(s, r, 8, f'=IF(D{r}="","",IFERROR(INDEX({GW},{m}),""))', "out")
    cell(s, r, 9, f'=IF(OR(G{r}="",F{r}=""),"",G{r}+F{r})', "out", fmt="0.00")
    cell(s, r, 10, f'=IF(I{r}="","",ROUND(IF(H{r}="Có",IF($C$4="Phần nhóm",MIN(7,G{r})+F{r},MIN(8,I{r})),I{r})*4,0)/4)',
         "out", fmt="0.00")
    cell(s, r, 11, f'=IF(J{r}="","",IF(J{r}>=9,"Xuất sắc",IF(J{r}>=8,"Giỏi",IF(J{r}>=6.5,"Khá",IF(J{r}>=5,"Trung bình",IF(J{r}>=3,"Yếu","Kém"))))))', "out")
    cell(s, r, 12, f'=IF(B{r}="","",IF(D{r}="","Chưa chọn nhóm",IF(ISNA(MATCH(D{r},{GA},0)),"Mã nhóm không có",'
                   f'IF(F{r}="","Chưa nhập D",IF(COUNTIF($B${s1}:$B${s2},B{r})>1,"MSSV bị trùng","OK")))))', "out", align=wrap)
    cell(s, r, 13, None, "in", align=wrap)
s.conditional_formatting.add(f"L{s1}:L{s2}", FormulaRule(formula=[f'AND(L{s1}<>"",L{s1}<>"OK")'],
                             fill=PatternFill("solid", fgColor="F8CBAD")))
s.conditional_formatting.add(f"J{s1}:J{s2}", CellIsRule(operator="lessThan", formula=["5"],
                             fill=PatternFill("solid", fgColor="F8CBAD")))

# ================= Tổng hợp =================
t = wb.create_sheet("Tổng hợp")
t["A1"] = "TỔNG HỢP ĐIỂM THI VẤN ĐÁP"; t["A1"].font = title
t.column_dimensions["A"].width = 34; t.column_dimensions["B"].width = 14; t.column_dimensions["C"].width = 14
J = f"'Chấm cá nhân'!$J${s1}:$J${s2}"
K = f"'Chấm cá nhân'!$K${s1}:$K${s2}"
stats = [("Số sinh viên đã có điểm", f"=COUNT({J})", "0"),
         ("Điểm trung bình", f'=IF(COUNT({J})=0,"",AVERAGE({J}))', "0.00"),
         ("Điểm cao nhất", f'=IF(COUNT({J})=0,"",MAX({J}))', "0.00"),
         ("Điểm thấp nhất", f'=IF(COUNT({J})=0,"",MIN({J}))', "0.00"),
         ("Số nhóm đã chấm", f"=COUNT('Chấm nhóm'!$V${first}:$V${last})", "0"),
         ("Số nhóm đã nhận câu vớt", f"=COUNTIF('Chấm nhóm'!$W${first}:$W${last},\"Có\")", "0")]
hdr(t, 3, ["Chỉ số", "Giá trị"])
for i, (a, f, fmt) in enumerate(stats, 4):
    cell(t, i, 1, a, align=wrap); cell(t, i, 2, f, "out", fmt=fmt)
r0 = 4 + len(stats) + 1
hdr(t, r0, ["Xếp loại", "Số SV", "Tỷ lệ"])
for k, lab in enumerate(["Xuất sắc", "Giỏi", "Khá", "Trung bình", "Yếu", "Kém"], r0 + 1):
    cell(t, k, 1, lab, align=wrap)
    cell(t, k, 2, f'=COUNTIF({K},"{lab}")', "out", fmt="0")
    cell(t, k, 3, f'=IF($B$4=0,"",B{k}/$B$4)', "out", fmt="0%")

# ================= Ví dụ =================
v = wb.create_sheet("Ví dụ")
v["A1"] = "4 VÍ DỤ TRONG RUBRIC (cùng công thức với hai sheet chấm)"; v["A1"].font = title
hdr(v, 3, ["Ví dụ", "A", "B", "C1", "C2", "Câu vớt", "Điểm vớt", "Vớt thay cho", "C sau vớt", "Điểm nhóm", "D",
           "Điểm cuối (trần: Điểm cuối)", "Điểm cuối (trần: Phần nhóm)", "Kỳ vọng (rubric)"],
    [30, 7, 7, 7, 7, 13, 9, 11, 10, 10, 7, 14, 14, 16])
ex = [("1. Không dùng câu vớt", 2.5, 1.75, 2, 1.5, "Không cần", None, 1, 8.75),
      ("2. Dùng câu vớt, trả lời tốt", 3, 2, 2, 0, "Đã hỏi", 2, 1, 8),
      ("3. Nhóm từ chối câu vớt", 3, 2, 2, 0.5, "Nhóm từ chối", None, 1, 8.5),
      ("4. Dùng câu vớt, nhóm dưới trần", 2, 1.5, 0.5, 1, "Đã hỏi", 1.5, 0.75, 6.75)]
for i, e in enumerate(ex, 4):
    r = i
    for j, val in enumerate(e[:7], 1):
        cell(v, r, j, val, "in" if j > 1 else "f", align=wrap if j == 1 else center)
    cell(v, r, 8, f'=IF(F{r}="Đã hỏi",IF(MIN(N(D{r}),N(E{r}))<=0.5,IF(N(D{r})<=N(E{r}),"C1","C2"),""),"")', "out")
    cell(v, r, 9, f'=IF(H{r}="C1",MAX(N(D{r}),N(G{r})),N(D{r}))+IF(H{r}="C2",MAX(N(E{r}),N(G{r})),N(E{r}))', "out", fmt="0.00")
    cell(v, r, 10, f"=B{r}+C{r}+I{r}", "out", fmt="0.00")
    cell(v, r, 11, e[7], "in", fmt="0.00")
    cell(v, r, 12, f'=ROUND(IF(F{r}="Đã hỏi",MIN(8,J{r}+K{r}),J{r}+K{r})*4,0)/4', "out", fmt="0.00")
    cell(v, r, 13, f'=ROUND(IF(F{r}="Đã hỏi",MIN(7,J{r})+K{r},J{r}+K{r})*4,0)/4', "out", fmt="0.00")
    cell(v, r, 14, e[8], fmt="0.00")

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
